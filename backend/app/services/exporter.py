"""DOCX export of the approved Minutes of Meeting (plan §7, python-docx).

Thai is a complex script: Word picks the *complex-script* font (w:cs) and
size (w:szCs) for Thai characters, so both must be set on every run, not only
the ASCII font. Otherwise Thai text falls back to an arbitrary font/size.
"""
from __future__ import annotations

import io
from datetime import datetime, timezone
from zoneinfo import ZoneInfo

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from app.config import SUPPORTED_LANGUAGES, get_settings
from app.models import Meeting, MeetingReport

TOPIC_LABELS = {"resolved": "ได้ข้อสรุป", "mentioned": "มีการพูดถึง (ยังไม่มีข้อสรุป)",
                "unclear": "ไม่ชัดเจน", "not_mentioned": "ไม่ได้พูดถึง"}


def local(dt: datetime | None, fmt: str = "%d/%m/%Y %H:%M") -> str:
    if dt is None:
        return "-"
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(ZoneInfo(get_settings().display_timezone)).strftime(fmt)


def fmt_ms(ms: int | None) -> str:
    if ms is None:
        return "-"
    s = ms // 1000
    return f"{s // 3600:02d}:{s % 3600 // 60:02d}:{s % 60:02d}" if s >= 3600 else f"{s // 60:02d}:{s % 60:02d}"


def _evidence_ref(evidence: list[dict]) -> str:
    if not evidence:
        return "-"
    return ", ".join(fmt_ms(e["start_ms"]) for e in evidence[:4])


def _apply_font(run, font: str, size: float, bold: bool = False, color: RGBColor | None = None) -> None:  # noqa: ANN001
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), font)
    for tag in ("w:szCs",):
        el = rpr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            rpr.append(el)
        el.set(qn("w:val"), str(int(size * 2)))
    if bold:
        b_cs = rpr.find(qn("w:bCs"))
        if b_cs is None:
            rpr.append(OxmlElement("w:bCs"))
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:bidi"), "th-TH")


class _Writer:
    def __init__(self) -> None:
        self.font = get_settings().docx_font
        self.doc = Document()
        for section in self.doc.sections:
            section.left_margin = section.right_margin = Pt(60)
        normal = self.doc.styles["Normal"]
        normal.font.name = self.font
        normal.font.size = Pt(11)
        rpr = normal.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.insert(0, rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rfonts.set(qn(attr), self.font)

    def para(self, text: str = "", *, size: float = 11, bold: bool = False, color: RGBColor | None = None,
             style: str | None = None, align=None):  # noqa: ANN001
        p = self.doc.add_paragraph(style=style)
        if align is not None:
            p.alignment = align
        if text:
            _apply_font(p.add_run(text), self.font, size, bold, color)
        return p

    def heading(self, text: str, size: float = 14) -> None:
        p = self.para(text, size=size, bold=True, color=RGBColor(0x1F, 0x3A, 0x5F))
        p.paragraph_format.space_before = Pt(12)

    def table(self, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
        t = self.doc.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        for i, h in enumerate(headers):
            cell = t.rows[0].cells[i]
            cell.text = ""
            _apply_font(cell.paragraphs[0].add_run(h), self.font, 10.5, bold=True)
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:fill"), "E8EEF5")
            cell._element.get_or_add_tcPr().append(shd)
        for row in rows:
            cells = t.add_row().cells
            for i, value in enumerate(row):
                cells[i].text = ""
                _apply_font(cells[i].paragraphs[0].add_run(value), self.font, 10.5)
        if widths:
            for row in t.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Pt(w)


def build_docx(meeting: Meeting, report: MeetingReport, approver_name: str | None,
               segments: list[dict] | None = None) -> bytes:
    s = report.summary_json
    w = _Writer()
    w.para("รายงานการประชุม (Minutes of Meeting)", size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    w.para(s.get("meeting_title") or meeting.title, size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

    info = [
        ["วันที่สร้างรายการ", local(meeting.created_at)],
        ["ไฟล์ต้นทาง", meeting.source_filename or "-"],
        ["ภาษา", SUPPORTED_LANGUAGES.get(meeting.language, meeting.language)],
        ["ความยาวเสียง", fmt_ms((meeting.audio_info or {}).get("duration_ms"))],
        ["ฉบับที่", str(report.version)],
        ["ผู้อนุมัติ", f"{approver_name or '-'} ({local(report.approved_at)})" if report.approved_at else "ยังไม่อนุมัติ"],
    ]
    w.table(["รายการ", "รายละเอียด"], info, widths=[130, 330])

    w.heading("1. สรุปภาพรวม")
    w.para(s.get("summary") or "-")

    w.heading("2. สถานะหัวข้อ (Topic Checklist)")
    w.table(["หัวข้อ", "สถานะ", "ข้อสรุป", "อ้างอิงเวลา"],
            [[t["title"], TOPIC_LABELS.get(t["status"], t["status"]), t.get("resolution") or "-",
              _evidence_ref(t.get("evidence", []))] for t in s.get("topics", [])] or [["-", "-", "-", "-"]],
            widths=[130, 110, 160, 60])

    w.heading("3. มติที่ประชุม")
    if s.get("decisions"):
        for i, d in enumerate(s["decisions"], 1):
            w.para(f"{i}. {d['text']}  [{_evidence_ref(d.get('evidence', []))}]")
    else:
        w.para("ไม่มีมติที่ระบุได้จากการประชุม")

    w.heading("4. งานที่ต้องติดตาม (Action Items)")
    w.table(["งาน", "ผู้รับผิดชอบ", "กำหนดส่ง", "อ้างอิงเวลา"],
            [[a["task"], a.get("owner") or "ยังไม่ระบุ", a.get("due_date") or "ยังไม่ระบุ",
              _evidence_ref(a.get("evidence", []))] for a in s.get("action_items", [])] or [["-", "-", "-", "-"]],
            widths=[200, 100, 100, 60])

    if s.get("open_questions"):
        w.heading("5. ประเด็นที่ยังเปิดอยู่")
        for q in s["open_questions"]:
            w.para(f"• {q['text']}")

    if segments:
        w.heading("ภาคผนวก: บทถอดเสียง (Transcript)")
        for seg in segments:
            p = w.para()
            _apply_font(p.add_run(f"[{fmt_ms(seg['start_ms'])}] "), w.font, 9.5, color=RGBColor(0x66, 0x66, 0x66))
            _apply_font(p.add_run(seg["text"]), w.font, 10)

    w.para()
    w.para(f"เอกสารนี้สร้างโดยระบบ AI Meeting Minutes เมื่อ {local(datetime.now(timezone.utc))} "
           "เนื้อหาร่างโดย AI และผ่านการตรวจสอบโดยผู้อนุมัติ ไม่มีการแยกเสียงผู้พูด (speaker diarization)",
           size=9, color=RGBColor(0x88, 0x88, 0x88))
    buf = io.BytesIO()
    w.doc.save(buf)
    return buf.getvalue()


def report_filename(meeting: Meeting, report: MeetingReport) -> str:
    date = local(meeting.created_at, "%Y%m%d") if meeting.created_at else "meeting"
    return f"MoM_{date}_{meeting.title}_v{report.version}.docx"
